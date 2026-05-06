"""AI 資料格式整理工具 — 上傳原始資料與目標格式，Groq Llama 自動判斷結構並轉換"""

import io
import json
import os
import zipfile
from datetime import datetime

from groq import Groq
import pandas as pd
import streamlit as st

HISTORY_FILE = "format_history.json"
MODEL = "llama-3.3-70b-versatile"

st.set_page_config(page_title="AI 資料格式整理", page_icon="📊", layout="wide")

# ── persistence ───────────────────────────────────────────────────────────────

def load_history() -> list:
    if os.path.exists(HISTORY_FILE):
        try:
            with open(HISTORY_FILE, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            return []
    return []

def save_history(history: list) -> None:
    with open(HISTORY_FILE, "w", encoding="utf-8") as f:
        json.dump(history[:200], f, ensure_ascii=False, indent=2)

def append_log(entry: dict) -> None:
    history = load_history()
    if not history or history[0].get("ts") != entry.get("ts"):
        history.insert(0, entry)
        save_history(history)

# ── API client ────────────────────────────────────────────────────────────────

def _api_key_from_env() -> str | None:
    try:
        return st.secrets.get("GROQ_API_KEY")
    except Exception:
        pass
    return os.environ.get("GROQ_API_KEY")

@st.cache_resource
def _make_client(key: str) -> Groq:
    return Groq(api_key=key)

# ── file reading ──────────────────────────────────────────────────────────────

def read_raw_file(uploaded, sheet_name=None) -> pd.DataFrame | None:
    uploaded.seek(0)
    name = uploaded.name.lower()
    try:
        if name.endswith(".csv"):
            return pd.read_csv(uploaded)
        if name.endswith((".xlsx", ".xls")):
            engine = "xlrd" if name.endswith(".xls") else "openpyxl"
            return pd.read_excel(uploaded, sheet_name=sheet_name if sheet_name is not None else 0, engine=engine)
        st.error("原始資料請上傳 CSV 或 Excel 檔案")
    except Exception as e:
        st.error(f"讀取失敗：{e}")
    return None

def _excel_sheets(uploaded) -> list[str]:
    """回傳 Excel 各工作表名稱，不影響 seek 位置。"""
    uploaded.seek(0)
    try:
        raw = uploaded.read()
    finally:
        uploaded.seek(0)
    name = uploaded.name.lower()
    if not name.endswith((".xlsx", ".xls")):
        return []
    engine = "xlrd" if name.endswith(".xls") else "openpyxl"
    # 優先用 openpyxl 直接讀（不走 pandas 引擎偵測）
    if engine == "openpyxl":
        try:
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True)
            sheets = list(wb.sheetnames)
            wb.close()
            if sheets:
                return sheets
        except Exception:
            pass
    # fallback：pd.ExcelFile 加明確 engine
    try:
        return list(pd.ExcelFile(io.BytesIO(raw), engine=engine).sheet_names)
    except Exception:
        pass
    return []

def merge_raw_files(uploaded_files) -> tuple[pd.DataFrame | None, list[str]]:
    """合併多個原始資料檔案為一個 DataFrame，回傳 (merged_df, file_name_list)"""
    dfs, names = [], []
    for f in uploaded_files:
        df = read_raw_file(f)
        if df is not None:
            dfs.append(df)
            names.append(f.name)
    if not dfs:
        return None, []
    try:
        return pd.concat(dfs, ignore_index=True), names
    except Exception as e:
        st.error(f"合併失敗：{e}")
        return None, names

def read_target_file(uploaded, sheet_name=None) -> dict | None:
    """
    Returns dict:
      { "type": "excel" | "word",
        "df": DataFrame,          # for excel
        "tables": [...],          # for word
        "paragraphs": [...],      # for word
        "raw_bytes": bytes }
    """
    uploaded.seek(0)
    name = uploaded.name.lower()
    raw_bytes = uploaded.read()
    uploaded.seek(0)

    if name.endswith((".xlsx", ".xls")):
        try:
            engine = "xlrd" if name.endswith(".xls") else "openpyxl"
            df = pd.read_excel(io.BytesIO(raw_bytes), sheet_name=sheet_name if sheet_name is not None else 0, engine=engine)
            return {"type": "excel", "df": df, "raw_bytes": raw_bytes}
        except Exception as e:
            st.error(f"讀取 Excel 失敗：{e}")
            return None

    if name.endswith(".csv"):
        try:
            df = pd.read_csv(io.BytesIO(raw_bytes))
            return {"type": "excel", "df": df, "raw_bytes": raw_bytes}
        except Exception as e:
            st.error(f"讀取 CSV 失敗：{e}")
            return None

    if name.endswith(".docx"):
        try:
            from docx import Document
            from docx.oxml.ns import qn as _qn
            doc = Document(io.BytesIO(raw_bytes))

            # 逐一掃描 body 子元素，取每個表格前最近的段落文字作為表格標題
            tables, paragraphs, table_titles = [], [], []
            last_para_text = ""
            for child in doc.element.body:
                tag = child.tag.split("}")[-1] if "}" in child.tag else child.tag
                if tag == "p":
                    text = "".join(n.text or "" for n in child.iter() if n.tag.split("}")[-1] == "t").strip()
                    if text:
                        last_para_text = text
                        paragraphs.append({"style": tag, "text": text[:200]})
                elif tag == "tbl":
                    rows = []
                    for tr in child.findall(_qn("w:tr")):
                        row = [
                            "".join(t.text or "" for t in tc.iter(_qn("w:t"))).strip()
                            for tc in tr.findall(_qn("w:tc"))
                        ]
                        rows.append(row)
                    tables.append(rows)
                    table_titles.append(last_para_text)

            return {
                "type": "word",
                "tables": tables,
                "paragraphs": paragraphs,
                "table_titles": table_titles,
                "raw_bytes": raw_bytes,
            }
        except Exception as e:
            st.error(f"讀取 Word 失敗：{e}")
            return None

    st.error("目標格式請上傳 CSV、Excel (.xlsx) 或 Word (.docx)")
    return None

# ── describe structure ────────────────────────────────────────────────────────

def describe_raw_df(df: pd.DataFrame, n_sample: int = 5) -> str:
    lines = []
    for col in df.columns:
        series = df[col].dropna()
        dtype = str(df[col].dtype)
        unique_count = series.nunique()
        if unique_count <= 12:
            lines.append(f"  {col!r} ({dtype}) — 唯一值: {series.unique().tolist()}")
        else:
            lines.append(f"  {col!r} ({dtype}) — 範例: {series.head(n_sample).tolist()}，共 {unique_count} 種")
    return "\n".join(lines)

def describe_target(target: dict) -> str:
    if target["type"] == "excel":
        df = target["df"]
        lines = [f"格式類型：Excel / CSV，共 {len(df.columns)} 欄"]
        for col in df.columns:
            series = df[col].dropna()
            if series.empty:
                lines.append(f"  {col!r} — 無範例資料")
            elif series.nunique() <= 10:
                lines.append(f"  {col!r} — 唯一值: {series.unique().tolist()}")
            else:
                lines.append(f"  {col!r} — 範例: {series.head(3).tolist()}")
        return "\n".join(lines)

    lines = ["格式類型：Word (.docx)"]
    for i, tbl in enumerate(target.get("tables", [])):
        headers = tbl[0] if tbl else []
        sample = tbl[1:4] if len(tbl) > 1 else []
        lines.append(f"  表格 {i+1}：欄位 {headers}，共 {len(tbl)} 列（含標題），範例資料：{sample}")
    for p in target.get("paragraphs", [])[:10]:
        lines.append(f"  段落（{p['style']}）：{p['text'][:80]}")
    return "\n".join(lines)

# ── AI Mapping ─────────────────────────────────────────────────────────────────

_SYSTEM_PROMPT = """你是資料結構分析與轉換專家。任務：
1. 深度解讀「目標格式」的結構意圖（欄位語意、資料型態、值的模式、Word 表格版型等）
2. 將「原始資料」的欄位對應到目標格式，決定每個欄位需要的轉換方式

支援的轉換類型（transform 欄位）：
- "direct"      : 直接複製或改名
- "merge"       : 合併多欄位，join_sep 為分隔符
- "value_map"   : 值對應轉換（如 "M"→"男"、"Y"→"是"）
- "date_format" : 重新格式化日期，target_fmt 為 strftime 格式字串
- "number_fmt"  : 數字格式化（如加千分位、保留小數）
- "empty"       : 無對應欄位，輸出空白

只回傳合法 JSON，不要 markdown 包裹或說明文字。
對每一個目標欄位都要有一筆 mapping（source_cols 無對應則填空陣列）。"""


def _strip_json(text: str) -> str:
    text = text.strip()
    if "```" in text:
        parts = text.split("```")
        for i, part in enumerate(parts):
            if i % 2 == 1:
                text = part.lstrip("json").strip()
                break
    return text


def get_mapping(client: Groq, raw_df: pd.DataFrame, target: dict) -> dict:
    raw_desc = describe_raw_df(raw_df)
    raw_sample = raw_df.head(5).to_string(index=False)
    tgt_desc = describe_target(target)

    resp = client.chat.completions.create(
        model=MODEL,
        max_tokens=4096,
        temperature=0.2,
        messages=[
            {"role": "system", "content": _SYSTEM_PROMPT},
            {"role": "user", "content": f"""=== 原始資料 ===
欄位結構：
{raw_desc}

前 5 筆範例：
{raw_sample}

=== 目標格式（使用者上傳的範例，請分析其結構意圖）===
{tgt_desc}

請輸出轉換計畫（JSON）：
{{
  "structure_analysis": "你對目標格式結構的解讀（2–3 句）",
  "output_type": "excel 或 word",
  "mappings": [
    {{
      "target_col": "目標欄位或表格標題名稱",
      "source_cols": ["原始欄位名"],
      "transform": "direct | merge | value_map | date_format | number_fmt | empty",
      "join_sep": " ",
      "value_map": {{"原始值": "目標值"}},
      "target_fmt": "%Y/%m/%d",
      "decimal_places": 2,
      "note": "說明或 null"
    }}
  ]
}}"""},
        ],
    )

    text = _strip_json(resp.choices[0].message.content)
    return json.loads(text)

# ── Apply mapping ──────────────────────────────────────────────────────────────

def _safe_str(v) -> str:
    if pd.isna(v):
        return ""
    return str(v)

def apply_mapping_to_df(raw_df: pd.DataFrame,
                         target_cols: list,
                         mapping: dict) -> pd.DataFrame:
    result = pd.DataFrame(index=range(len(raw_df)), columns=target_cols)

    for m in mapping.get("mappings", []):
        target = m.get("target_col")
        if target not in target_cols:
            continue
        transform = m.get("transform", "direct")
        sources = [s for s in m.get("source_cols", []) if s in raw_df.columns]

        if transform == "empty" or not sources:
            result[target] = None
            continue

        if transform == "merge":
            sep = m.get("join_sep", " ")
            result[target] = raw_df[sources].apply(
                lambda row: sep.join(_safe_str(v) for v in row if _safe_str(v)),
                axis=1,
            ).values

        elif transform == "value_map":
            vmap = m.get("value_map", {})
            result[target] = raw_df[sources[0]].map(
                lambda v: vmap.get(_safe_str(v), _safe_str(v))
            ).values

        elif transform == "date_format":
            fmt = m.get("target_fmt", "%Y/%m/%d")
            def _fmt_date(v, f=fmt):
                if pd.isna(v):
                    return None
                try:
                    return pd.to_datetime(v).strftime(f)
                except Exception:
                    return _safe_str(v)
            result[target] = raw_df[sources[0]].map(_fmt_date).values

        elif transform == "number_fmt":
            dp = m.get("decimal_places", 2)
            def _fmt_num(v, d=dp):
                try:
                    return f"{float(v):,.{d}f}"
                except Exception:
                    return _safe_str(v)
            result[target] = raw_df[sources[0]].map(_fmt_num).values

        else:  # direct
            result[target] = raw_df[sources[0]].values

    return result

# ── Output generators ──────────────────────────────────────────────────────────

def generate_excel(result_df: pd.DataFrame) -> bytes:
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        result_df.to_excel(writer, index=False, sheet_name="整理結果")
    return buf.getvalue()

def _fill_word_table(table, result_df: pd.DataFrame) -> None:
    """將 result_df 的資料列填入 Word table（清除舊資料列後重新寫入）"""
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    import copy

    header_cells = [cell.text.strip() for cell in table.rows[0].cells]
    while len(table.rows) > 1:
        table._tbl.remove(table.rows[-1]._tr)

    for _, row_data in result_df.iterrows():
        new_tr = copy.deepcopy(table.rows[0]._tr)
        new_row_cells = new_tr.findall(qn("w:tc"))
        for j, header in enumerate(header_cells):
            if j < len(new_row_cells):
                tc = new_row_cells[j]
                for p in tc.findall(qn("w:p")):
                    tc.remove(p)
                p_elem = OxmlElement("w:p")
                r_elem = OxmlElement("w:r")
                t_elem = OxmlElement("w:t")
                val = row_data.get(header, "")
                t_elem.text = "" if pd.isna(val) else str(val)
                r_elem.append(t_elem)
                p_elem.append(r_elem)
                tc.append(p_elem)
        table._tbl.append(new_tr)


def generate_word(result_df: pd.DataFrame,
                  target: dict,
                  mapping: dict,
                  target_table_idxs: list[int] | None = None) -> bytes:
    from docx import Document

    raw_bytes = target.get("raw_bytes", b"")
    doc = Document(io.BytesIO(raw_bytes)) if raw_bytes else Document()

    if doc.tables:
        idxs = target_table_idxs if target_table_idxs else [0]
        for idx in idxs:
            idx = min(idx, len(doc.tables) - 1)
            _fill_word_table(doc.tables[idx], result_df)
    else:
        table = doc.add_table(rows=1, cols=len(result_df.columns))
        table.style = "Table Grid"
        for j, col in enumerate(result_df.columns):
            table.rows[0].cells[j].text = col
        for _, row_data in result_df.iterrows():
            row = table.add_row()
            for j, col in enumerate(result_df.columns):
                val = row_data[col]
                row.cells[j].text = "" if pd.isna(val) else str(val)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()

def generate_zip(results: list) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for item in results:
            zf.writestr(item["filename"], item["data"])
    return buf.getvalue()

# ── Shared processing logic ────────────────────────────────────────────────────

def run_conversion(client: Groq,
                   raw_df: pd.DataFrame,
                   target: dict,
                   raw_name: str = "",
                   force_table_idxs: list[int] | None = None) -> dict:
    """
    執行一次完整的 AI 分析與轉換，回傳 dict 包含所有結果。
    force_table_idxs：指定要填入資料的表格索引清單（0-based）；
                      以第一個索引的欄位做 AI 映射，所有指定表格都會填入相同資料。
                      None 或空 list 則 fallback 為欄位最多的表格。
    """
    target_type = target["type"]
    target_table_idxs: list[int] = []

    if target_type == "excel":
        target_cols = list(target["df"].columns)
    else:
        tables = target.get("tables", [])
        if tables:
            if force_table_idxs:
                # 使用者明確選取（可能多個）
                target_table_idxs = [min(i, len(tables) - 1) for i in force_table_idxs]
                first = target_table_idxs[0]
                target_cols = list(tables[first][0]) if tables[first] else []
            else:
                # fallback：欄位最多的表格
                idx = max(range(len(tables)), key=lambda i: len(tables[i][0]) if tables[i] else 0)
                target_table_idxs = [idx]
                target_cols = list(tables[idx][0])
        else:
            target_cols = []

    mapping = get_mapping(client, raw_df, target)

    result_df = apply_mapping_to_df(raw_df, target_cols, mapping)

    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    prefix = raw_name.rsplit(".", 1)[0] if raw_name else "result"

    if target_type == "word":
        try:
            out_bytes = generate_word(result_df, target, mapping, target_table_idxs)
            out_name = f"{prefix}_{ts}.docx"
            mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        except Exception:
            out_bytes = generate_excel(result_df)
            out_name = f"{prefix}_{ts}.xlsx"
            mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    else:
        out_bytes = generate_excel(result_df)
        out_name = f"{prefix}_{ts}.xlsx"
        mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    matched = sum(
        1 for m in mapping.get("mappings", [])
        if any(s in raw_df.columns for s in m.get("source_cols", []))
    )

    return {
        "result_df": result_df,
        "mapping": mapping,
        "out_bytes": out_bytes,
        "out_name": out_name,
        "mime": mime,
        "matched": matched,
        "target_cols": target_cols,
        "target_type": target_type,
    }

# ── Mapping display helper ─────────────────────────────────────────────────────

_LABELS = {
    "direct": "直接複製",
    "merge": "合併欄位",
    "value_map": "值對應轉換",
    "date_format": "日期格式轉換",
    "number_fmt": "數字格式化",
    "empty": "無對應（空白）",
}

def show_mapping_table(mapping: dict, raw_df: pd.DataFrame, target_cols: list) -> None:
    map_rows, matched, unmatched = [], 0, 0
    for m in mapping.get("mappings", []):
        sources = [s for s in m.get("source_cols", []) if s in raw_df.columns]
        transform = m.get("transform", "direct")
        if sources:
            matched += 1
            src_str = " + ".join(sources)
        else:
            unmatched += 1
            src_str = "⚠️ 無對應"
        extra = ""
        if transform == "value_map" and m.get("value_map"):
            extra = f"  {m['value_map']}"
        elif transform == "date_format" and m.get("target_fmt"):
            extra = f"  → {m['target_fmt']}"
        elif transform == "merge" and m.get("join_sep") not in (None, " ", ""):
            extra = f"  分隔符：{m['join_sep']!r}"
        map_rows.append({
            "目標欄位": m.get("target_col", ""),
            "來源欄位": src_str,
            "轉換方式": _LABELS.get(transform, transform) + extra,
            "備註": m.get("note") or "",
        })

    mc1, mc2, mc3 = st.columns(3)
    mc1.metric("目標欄位總數", len(target_cols))
    mc2.metric("成功映射", matched)
    mc3.metric("無對應來源", unmatched)

    st.dataframe(
        pd.DataFrame(map_rows),
        use_container_width=True,
        hide_index=True,
        column_config={
            "目標欄位": st.column_config.TextColumn(width="small"),
            "來源欄位": st.column_config.TextColumn(width="medium"),
            "轉換方式": st.column_config.TextColumn(width="medium"),
            "備註":     st.column_config.TextColumn(width="large"),
        },
    )

# ── Sidebar ────────────────────────────────────────────────────────────────────

with st.sidebar:
    st.markdown("## ⚙️ 設定")
    st.divider()

    env_key = _api_key_from_env()
    if env_key:
        api_key = env_key
        st.success("✅ Groq Key 已從環境載入")
    else:
        api_key = st.text_input(
            "Groq API Key（免費）",
            type="password",
            placeholder="gsk_...",
            help="免費申請：console.groq.com，用 Google 帳號即可",
        )
        if api_key:
            st.success("✅ Groq Key 已輸入")
    st.caption("📌 [免費取得 Groq Key](https://console.groq.com)")

    st.divider()
    st.markdown("### 📖 使用說明")
    st.markdown("""
**① 上傳原始資料**
支援上傳**多個** Excel / CSV，系統會自動合併成一份再整理。

**② 上傳目標格式**
你希望整理成的格式範例：
- Excel / CSV → 輸出 Excel
- Word (.docx) → 輸出 Word

**③ 點擊「開始整理」**
AI 分析結構、映射欄位、轉換資料，完成後即可下載。

---

**📦 批次整理**：設定多組「原始資料＋格式」，一次處理後打包成 ZIP 下載。
    """)

    st.divider()
    history_sidebar = load_history()
    st.metric("累計操作次數", len(history_sidebar))
    if history_sidebar:
        last = history_sidebar[0]
        st.caption(f"最近：{last.get('date','')} {last.get('time','')}")

    st.divider()
    st.caption(f"模型：{MODEL}（Groq 免費）")
    st.caption("原始資料支援：CSV、xlsx、xls")
    st.caption("目標格式支援：CSV、xlsx、docx")

# ── Main ───────────────────────────────────────────────────────────────────────

st.title("📊 AI 資料格式整理工具")
st.caption("上傳原始資料 + 你想要的格式範例，AI 自動判斷結構、映射欄位並輸出整理後的檔案")

tab_main, tab_batch, tab_history = st.tabs(["📁 單一整理", "📦 批次整理", "📋 操作歷史"])

# ══════════════════════════════════════════════════════════════════════════════
# Tab 1: 單一整理（支援多個原始資料自動合併）
# ══════════════════════════════════════════════════════════════════════════════
with tab_main:
    col_l, col_r = st.columns(2, gap="large")

    with col_l:
        st.subheader("① 原始資料")
        st.caption("可同時選取多個 Excel / CSV，系統自動合併（欄位相同效果最佳）")
        raw_files = st.file_uploader(
            "拖曳或點擊上傳原始資料",
            type=["csv", "xlsx", "xls"],
            key="raw_upload",
            label_visibility="collapsed",
            accept_multiple_files=True,
        )

        raw_df = None
        if raw_files:
            if len(raw_files) == 1:
                rf0 = raw_files[0]
                _raw_sheet = None
                if rf0.name.lower().endswith((".xlsx", ".xls")):
                    _rs = _excel_sheets(rf0)
                    if len(_rs) > 1:
                        _raw_sheet = st.selectbox("工作表（原始資料）", _rs, key="raw_sheet_single")
                raw_df = read_raw_file(rf0, sheet_name=_raw_sheet)
                if raw_df is not None:
                    st.success(f"**{rf0.name}** — {len(raw_df):,} 筆 × {len(raw_df.columns)} 欄")
            else:
                raw_df, merged_names = merge_raw_files(raw_files)
                if raw_df is not None:
                    st.success(f"已合併 {len(merged_names)} 個檔案 → **{len(raw_df):,} 筆 × {len(raw_df.columns)} 欄**")
                    with st.expander(f"合併的檔案清單（{len(merged_names)} 個）"):
                        for n in merged_names:
                            st.write(f"• {n}")
            if raw_df is not None:
                st.dataframe(raw_df.head(5), use_container_width=True, height=200)
                with st.expander("所有欄位名稱"):
                    st.write(list(raw_df.columns))

    with col_r:
        st.subheader("② 想整理成的格式（上傳範例）")
        st.caption("副檔名決定輸出格式：`.xlsx` → Excel，`.docx` → Word")
        tmpl_file = st.file_uploader(
            "拖曳或點擊上傳目標格式",
            type=["csv", "xlsx", "docx"],
            key="tmpl_upload",
            label_visibility="collapsed",
        )

        target = None
        if tmpl_file:
            _tmpl_sheet = None
            if tmpl_file.name.lower().endswith((".xlsx", ".xls")):
                _ts = _excel_sheets(tmpl_file)
                if len(_ts) > 1:
                    _tmpl_sheet = st.selectbox("工作表（目標格式）", _ts, key="tmpl_sheet_single")
            target = read_target_file(tmpl_file, sheet_name=_tmpl_sheet)
            if target is not None:
                if target["type"] == "excel":
                    df_preview = target["df"]
                    st.success(f"**{tmpl_file.name}** (Excel) — {len(df_preview.columns)} 欄，{len(df_preview)} 列範例")
                    st.dataframe(df_preview.head(5), use_container_width=True, height=200)
                    with st.expander("所有目標欄位"):
                        st.write(list(df_preview.columns))
                else:
                    tables = target.get("tables", [])
                    paras = target.get("paragraphs", [])
                    st.success(f"**{tmpl_file.name}** (Word) — {len(tables)} 個表格，{len(paras)} 個段落")

    # Word 表格選擇器（每個原始資料各自選擇，複選）
    word_options = []   # 供 button handler 取用
    if target is not None and target["type"] == "word" and raw_files:
        w_tables = target.get("tables", [])
        w_titles = target.get("table_titles", [])

        if w_tables:
            def _table_label(i, tbl, title=""):
                headers = tbl[0] if tbl else []
                cols_preview = " | ".join(str(h) for h in headers[:5])
                if len(headers) > 5:
                    cols_preview += " | …"
                name = title.strip() if title.strip() else f"表格 {i+1}"
                return f"{name}（{len(headers)} 欄）：{cols_preview}"

            word_options = [
                _table_label(i, tbl, w_titles[i] if i < len(w_titles) else "")
                for i, tbl in enumerate(w_tables)
            ]

            st.markdown("### ③ 設定每個原始資料要填入的表格")
            for fi, rf in enumerate(raw_files):
                with st.container(border=True):
                    c_name, c_sel = st.columns([2, 5])
                    with c_name:
                        st.markdown(f"**{rf.name}**")
                    with c_sel:
                        st.multiselect(
                            "填入表格",
                            options=word_options,
                            key=f"single_file_tables_{fi}",
                            label_visibility="collapsed",
                            help="可複選；AI 以第一個選取的表格欄位做映射",
                        )
                        sel_preview = st.session_state.get(f"single_file_tables_{fi}") or []
                        if sel_preview:
                            first_idx = word_options.index(sel_preview[0])
                            chosen = w_tables[first_idx]
                            if chosen and len(chosen) > 1:
                                st.dataframe(
                                    pd.DataFrame(chosen[1:4], columns=chosen[0]),
                                    use_container_width=True,
                                    height=130,
                                )

    st.divider()

    # 自訂輸出檔名
    single_custom_name = st.text_input(
        "輸出檔案名稱（選填，不含副檔名）",
        placeholder="預設使用原始資料檔名",
        key="single_output_name",
    )

    ready = bool(raw_files) and (tmpl_file is not None)
    can_run = ready and bool(api_key) and (target is not None)

    if not api_key:
        st.warning("⚠️ 請先在左側輸入 Groq API Key（免費申請：console.groq.com）")

    if st.button(
        "🚀 開始整理",
        type="primary",
        use_container_width=True,
        disabled=(not can_run),
        key="run_btn",
    ):
        client = _make_client(api_key)
        now = datetime.now()
        ts = now.strftime("%Y%m%d_%H%M%S")
        raw_names_str = ", ".join(f.name for f in raw_files)

        # ── Word 多原始資料：每個檔案分別 AI 分析，填入各自選取的表格 ──────────
        if target["type"] == "word" and len(raw_files) >= 1 and word_options:
            from docx import Document as _DocX
            doc_obj = _DocX(io.BytesIO(target["raw_bytes"]))
            w_tables_rt = target.get("tables", [])
            all_matched_info = []
            total_rows_processed = 0

            for fi, rf in enumerate(raw_files):
                sel_labels = st.session_state.get(f"single_file_tables_{fi}") or []
                if not sel_labels:
                    st.warning(f"⚠️ {rf.name} 未選取表格，跳過")
                    continue

                tbl_idxs = [word_options.index(s) for s in sel_labels]
                # 單一原始檔直接用 UI 階段已讀好的 raw_df，多檔才重新讀
                g_raw_df = raw_df if len(raw_files) == 1 else read_raw_file(rf)
                if g_raw_df is None:
                    continue

                # 以第一個選取表格的欄位為基準做 AI 映射
                first_tbl_headers = [
                    c.text.strip() for c in doc_obj.tables[tbl_idxs[0]].rows[0].cells
                ] if tbl_idxs[0] < len(doc_obj.tables) else []

                focused_target = {
                    "type": "excel",
                    "df": pd.DataFrame(columns=first_tbl_headers),
                    "raw_bytes": b"",
                }

                with st.spinner(f"AI 分析 {rf.name}…"):
                    try:
                        mapping = get_mapping(client, g_raw_df, focused_target)
                    except json.JSONDecodeError as e:
                        st.error(f"{rf.name}：AI 回傳格式錯誤（{e}），跳過")
                        continue
                    except Exception as e:
                        st.error(f"{rf.name}：AI 分析失敗（{e}），跳過")
                        continue

                result_df = apply_mapping_to_df(g_raw_df, first_tbl_headers, mapping)

                for idx in tbl_idxs:
                    if idx < len(doc_obj.tables):
                        _fill_word_table(doc_obj.tables[idx], result_df)

                matched = sum(
                    1 for m in mapping.get("mappings", [])
                    if any(s in g_raw_df.columns for s in m.get("source_cols", []))
                )
                total_rows_processed += len(g_raw_df)
                all_matched_info.append({
                    "file": rf.name, "mapping": mapping,
                    "result_df": result_df, "target_cols": first_tbl_headers,
                    "matched": matched, "tbl_idxs": tbl_idxs,
                })

                analysis = mapping.get("structure_analysis")
                if analysis:
                    st.info(f"**{rf.name}** AI 理解：{analysis}")
                with st.expander(f"📋 {rf.name} 轉換計畫"):
                    show_mapping_table(mapping, g_raw_df, first_tbl_headers)
                st.caption(f"**{rf.name}** → 填入表格 {[i+1 for i in tbl_idxs]}，{len(result_df):,} 筆")

            if not all_matched_info:
                st.error("所有原始資料均未處理（請確認已為每個檔案選取填入表格）")
                st.stop()

            buf = io.BytesIO()
            doc_obj.save(buf)
            out_bytes = buf.getvalue()
            out_name_base = single_custom_name.strip() or raw_files[0].name.rsplit(".", 1)[0]
            out_name = f"{out_name_base}_{ts}.docx"
            mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

            st.download_button(
                f"⬇️ 下載整理後的 Word（{out_name}）",
                data=out_bytes, file_name=out_name, mime=mime,
                use_container_width=True, type="primary",
            )
            for info in all_matched_info:
                append_log({
                    "ts": now.strftime("%Y-%m-%d %H:%M:%S"),
                    "date": now.strftime("%Y-%m-%d"),
                    "time": now.strftime("%H:%M:%S"),
                    "raw_file": info["file"],
                    "template_file": tmpl_file.name,
                    "output_file": out_name,
                    "output_type": "word",
                    "rows": len(info["result_df"]),
                    "mapped": info["matched"],
                    "total_cols": len(info["target_cols"]),
                })
            st.success(f"✅ 完成！共 {len(all_matched_info)} 個檔案，合計 {total_rows_processed:,} 筆")

        # ── Excel 或單純 Word 無選擇器：合併所有檔案後一次處理 ────────────────
        else:
            if len(raw_files) == 1:
                single_raw_df = raw_df  # UI 階段已用正確 sheet 讀取，直接使用
                raw_base = raw_files[0].name
            else:
                single_raw_df, _ = merge_raw_files(raw_files)
                raw_base = "merged"

            if single_raw_df is None:
                st.error("原始資料讀取失敗")
                st.stop()

            with st.spinner("AI 分析目標格式結構並規劃轉換方式中…"):
                try:
                    res = run_conversion(client, single_raw_df, target, raw_base)
                except json.JSONDecodeError as e:
                    st.error(f"AI 回傳格式錯誤，請再試一次。（{e}）")
                    st.stop()
                except Exception as e:
                    st.error(f"AI 分析失敗：{e}")
                    st.stop()

            if single_custom_name.strip():
                ext = res["out_name"].rsplit(".", 1)[-1]
                res["out_name"] = f"{single_custom_name.strip()}_{ts}.{ext}"

            analysis = res["mapping"].get("structure_analysis")
            if analysis:
                st.info(f"**AI 的結構理解：** {analysis}")

            st.subheader("轉換計畫")
            show_mapping_table(res["mapping"], single_raw_df, res["target_cols"])

            st.subheader("整理結果預覽")
            st.dataframe(res["result_df"].head(10), use_container_width=True)
            st.caption(f"共 {len(res['result_df']):,} 筆")

            st.download_button(
                f"⬇️ 下載整理後的{'Word' if res['target_type'] == 'word' else 'Excel'}（{res['out_name']}）",
                data=res["out_bytes"], file_name=res["out_name"], mime=res["mime"],
                use_container_width=True, type="primary",
            )
            append_log({
                "ts":            now.strftime("%Y-%m-%d %H:%M:%S"),
                "date":          now.strftime("%Y-%m-%d"),
                "time":          now.strftime("%H:%M:%S"),
                "raw_file":      raw_names_str,
                "template_file": tmpl_file.name,
                "output_file":   res["out_name"],
                "output_type":   res["target_type"],
                "rows":          len(single_raw_df),
                "mapped":        res["matched"],
                "total_cols":    len(res["target_cols"]),
            })
            st.success(f"✅ 完成！處理 {len(single_raw_df):,} 筆，映射 {res['matched']}/{len(res['target_cols'])} 個欄位")

    if bool(raw_files) and tmpl_file is None:
        st.info("👆 請上傳目標格式，「開始整理」按鈕就會啟用")
    elif not raw_files and tmpl_file is not None:
        st.info("👆 請上傳原始資料，「開始整理」按鈕就會啟用")

# ══════════════════════════════════════════════════════════════════════════════
# Tab 2: 批次整理（多組 原始資料 + 目標格式）
# ══════════════════════════════════════════════════════════════════════════════
with tab_batch:
    st.subheader("批次整理模式")
    st.caption("每組設定一份「原始資料（可多檔合併）」＋「目標格式」，一次處理多組，打包成 ZIP 下載")

    if "batch_group_ids" not in st.session_state:
        st.session_state.batch_group_ids = [0]
    if "batch_next_id" not in st.session_state:
        st.session_state.batch_next_id = 1

    st.write("")

    # ── render each group ──────────────────────────────────────────────────────
    for gid in list(st.session_state.batch_group_ids):
        group_idx = st.session_state.batch_group_ids.index(gid) + 1
        with st.expander(f"📂 群組 {group_idx}", expanded=True):
            bc1, bc2, bc_del = st.columns([5, 5, 1])

            with bc1:
                st.markdown("**原始資料**（可選多個，自動合併）")
                st.file_uploader(
                    "上傳原始資料",
                    type=["csv", "xlsx", "xls"],
                    key=f"batch_raw_{gid}",
                    label_visibility="collapsed",
                    accept_multiple_files=True,
                )
                b_raws = st.session_state.get(f"batch_raw_{gid}") or []
                if b_raws:
                    names_str = "、".join(f.name for f in b_raws)
                    st.caption(f"✅ {len(b_raws)} 個檔案：{names_str}")
                if len(b_raws) == 1 and b_raws[0].name.lower().endswith((".xlsx", ".xls")):
                    _brs = _excel_sheets(b_raws[0])
                    if len(_brs) > 1:
                        st.selectbox("工作表（原始資料）", _brs, key=f"batch_raw_sheet_{gid}")

            with bc2:
                st.markdown("**目標格式**")
                st.file_uploader(
                    "上傳目標格式",
                    type=["csv", "xlsx", "docx"],
                    key=f"batch_tmpl_{gid}",
                    label_visibility="collapsed",
                )
                b_tmpl = st.session_state.get(f"batch_tmpl_{gid}")
                if b_tmpl:
                    st.caption(f"✅ {b_tmpl.name}")
                if b_tmpl and b_tmpl.name.lower().endswith((".xlsx", ".xls")):
                    _bts = _excel_sheets(b_tmpl)
                    if len(_bts) > 1:
                        st.selectbox("工作表（目標格式）", _bts, key=f"batch_tmpl_sheet_{gid}")

            with bc_del:
                st.write("")
                st.write("")
                if len(st.session_state.batch_group_ids) > 1:
                    if st.button("❌", key=f"del_{gid}", help="移除此群組"):
                        st.session_state.batch_group_ids.remove(gid)
                        st.rerun()

            # Word 表格選擇器（上傳 docx 後才顯示，複選）
            if b_tmpl and b_tmpl.name.lower().endswith(".docx"):
                try:
                    _bt = read_target_file(b_tmpl)
                    _w_tables = _bt.get("tables", []) if _bt else []
                    _w_titles = _bt.get("table_titles", []) if _bt else []
                    if _w_tables:
                        def _blabel(i, tbl, title=""):
                            h = tbl[0] if tbl else []
                            prev = " | ".join(str(x) for x in h[:6])
                            if len(h) > 6:
                                prev += " | …"
                            name = title.strip() if title.strip() else f"表格 {i+1}"
                            return f"{name}（{len(h)} 欄）：{prev}"
                        _opts = [
                            _blabel(i, t, _w_titles[i] if i < len(_w_titles) else "")
                            for i, t in enumerate(_w_tables)
                        ]
                        st.multiselect(
                            "🎯 選擇要填入資料的表格（可複選）",
                            options=_opts,
                            key=f"batch_word_tbl_{gid}",
                            help="可同時選多個結構相同的表格，AI 以第一個選取的欄位做映射",
                        )
                except Exception:
                    pass

            # 自訂輸出檔名（每組各自設定）
            st.text_input(
                "輸出檔案名稱（選填，不含副檔名）",
                placeholder="預設使用原始資料檔名",
                key=f"batch_outname_{gid}",
                label_visibility="visible",
            )

    # ── add group button ───────────────────────────────────────────────────────
    if st.button("➕ 新增群組", use_container_width=False):
        new_id = st.session_state.batch_next_id
        st.session_state.batch_group_ids.append(new_id)
        st.session_state.batch_next_id += 1
        st.rerun()

    st.divider()

    # ── readiness check ────────────────────────────────────────────────────────
    batch_ready_ids = [
        gid for gid in st.session_state.batch_group_ids
        if (st.session_state.get(f"batch_raw_{gid}") or [])
        and st.session_state.get(f"batch_tmpl_{gid}")
    ]
    n_total = len(st.session_state.batch_group_ids)
    n_ready = len(batch_ready_ids)

    if n_ready > 0:
        st.info(f"已就緒 **{n_ready}/{n_total}** 組，點擊下方按鈕開始批次處理")
    else:
        st.info("請為每組上傳原始資料與目標格式")

    if not api_key:
        st.warning("⚠️ 請先在左側輸入 Groq API Key")

    if st.button(
        f"🚀 開始批次整理（{n_ready} 組）",
        type="primary",
        use_container_width=True,
        disabled=(n_ready == 0 or not api_key),
        key="batch_run_btn",
    ):
        client = _make_client(api_key)
        batch_results = []
        now = datetime.now()
        ts = now.strftime("%Y%m%d_%H%M%S")

        progress = st.progress(0, text="準備中…")

        for i, gid in enumerate(batch_ready_ids):
            group_label = f"群組 {st.session_state.batch_group_ids.index(gid) + 1}"
            progress.progress(i / n_ready, text=f"處理 {group_label}（{i+1}/{n_ready}）…")

            b_raws = st.session_state.get(f"batch_raw_{gid}") or []
            b_tmpl = st.session_state.get(f"batch_tmpl_{gid}")

            with st.container(border=True):
                st.markdown(f"**⚙️ {group_label}**")
                try:
                    # 讀取並合併原始資料
                    if len(b_raws) == 1:
                        _g_raw_sheet = st.session_state.get(f"batch_raw_sheet_{gid}")
                        g_raw_df = read_raw_file(b_raws[0], sheet_name=_g_raw_sheet)
                        g_raw_name = b_raws[0].name
                    else:
                        g_raw_df, g_merged_names = merge_raw_files(b_raws)
                        g_raw_name = f"merged_group{gid}"

                    if g_raw_df is None:
                        st.error("原始資料讀取失敗，跳過此群組")
                        continue

                    # 讀取目標格式
                    _g_tmpl_sheet = st.session_state.get(f"batch_tmpl_sheet_{gid}")
                    g_target = read_target_file(b_tmpl, sheet_name=_g_tmpl_sheet)
                    if g_target is None:
                        st.error("目標格式讀取失敗，跳過此群組")
                        continue

                    # 取得使用者選的表格索引清單（Word 才有，複選）
                    g_force_tbls = None
                    if g_target["type"] == "word":
                        sel_labels = st.session_state.get(f"batch_word_tbl_{gid}") or []
                        if sel_labels:
                            _rt_tables = g_target.get("tables", [])
                            _rt_titles = g_target.get("table_titles", [])
                            def _blabel_rt(i, tbl, title=""):
                                h = tbl[0] if tbl else []
                                prev = " | ".join(str(x) for x in h[:6])
                                if len(h) > 6:
                                    prev += " | …"
                                name = title.strip() if title.strip() else f"表格 {i+1}"
                                return f"{name}（{len(h)} 欄）：{prev}"
                            _opts_rt = [
                                _blabel_rt(i, t, _rt_titles[i] if i < len(_rt_titles) else "")
                                for i, t in enumerate(_rt_tables)
                            ]
                            g_force_tbls = []
                            for sl in sel_labels:
                                try:
                                    g_force_tbls.append(_opts_rt.index(sl))
                                except ValueError:
                                    pass
                            if not g_force_tbls:
                                g_force_tbls = None

                    # AI 分析與轉換
                    with st.spinner(f"AI 分析 {group_label}…"):
                        res = run_conversion(client, g_raw_df, g_target, g_raw_name, force_table_idxs=g_force_tbls)

                    analysis = res["mapping"].get("structure_analysis")
                    if analysis:
                        st.info(f"AI 結構理解：{analysis}")

                    # 套用自訂檔名
                    batch_custom_name = (st.session_state.get(f"batch_outname_{gid}") or "").strip()
                    if batch_custom_name:
                        ext = res["out_name"].rsplit(".", 1)[-1]
                        batch_ts = now.strftime("%Y%m%d_%H%M%S")
                        res["out_name"] = f"{batch_custom_name}_{batch_ts}.{ext}"

                    col_stat1, col_stat2 = st.columns(2)
                    col_stat1.metric("處理筆數", f"{len(res['result_df']):,}")
                    col_stat2.metric("映射欄位", f"{res['matched']}/{len(res['target_cols'])}")

                    st.dataframe(res["result_df"].head(5), use_container_width=True, height=160)
                    st.caption(f"輸出檔名：{res['out_name']}")

                    batch_results.append({
                        "filename": res["out_name"],
                        "data": res["out_bytes"],
                    })

                    log_entry = {
                        "ts":            now.strftime("%Y-%m-%d %H:%M:%S"),
                        "date":          now.strftime("%Y-%m-%d"),
                        "time":          now.strftime("%H:%M:%S"),
                        "raw_file":      ", ".join(f.name for f in b_raws),
                        "template_file": b_tmpl.name,
                        "output_file":   res["out_name"],
                        "output_type":   res["target_type"],
                        "rows":          len(g_raw_df),
                        "mapped":        res["matched"],
                        "total_cols":    len(res["target_cols"]),
                    }
                    append_log(log_entry)
                    st.success(f"✅ 完成")

                except json.JSONDecodeError as e:
                    st.error(f"AI 回傳格式錯誤：{e}")
                except Exception as e:
                    st.error(f"處理失敗：{e}")

        progress.progress(1.0, text="全部完成！")

        # ── download ───────────────────────────────────────────────────────────
        if batch_results:
            st.divider()
            if len(batch_results) == 1:
                item = batch_results[0]
                ext = item["filename"].rsplit(".", 1)[-1]
                mime = (
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                    if ext == "docx"
                    else "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                )
                st.download_button(
                    f"⬇️ 下載結果（{item['filename']}）",
                    data=item["data"],
                    file_name=item["filename"],
                    mime=mime,
                    use_container_width=True,
                    type="primary",
                )
            else:
                zip_bytes = generate_zip(batch_results)
                st.download_button(
                    f"⬇️ 下載全部結果（ZIP，{len(batch_results)} 個檔案）",
                    data=zip_bytes,
                    file_name=f"batch_results_{ts}.zip",
                    mime="application/zip",
                    use_container_width=True,
                    type="primary",
                )
            st.success(f"✅ 批次完成！成功處理 {len(batch_results)}/{n_ready} 組")
        else:
            st.error("所有群組均處理失敗")

# ══════════════════════════════════════════════════════════════════════════════
# Tab 3: 操作歷史
# ══════════════════════════════════════════════════════════════════════════════
with tab_history:
    st.subheader("操作歷史記錄")
    history = load_history()

    if not history:
        st.info("📭 尚無記錄。完成第一次整理後會自動記錄。")
    else:
        total_rows = sum(h.get("rows", 0) for h in history)
        hc1, hc2, hc3 = st.columns(3)
        hc1.metric("總操作次數", len(history))
        hc2.metric("累計處理筆數", f"{total_rows:,}")
        last = history[0]
        hc3.metric("最近操作", f"{last.get('date','')} {last.get('time','')}")

        st.write("")
        hist_df = pd.DataFrame(history).rename(columns={
            "date":          "日期",
            "time":          "時間",
            "raw_file":      "原始檔案",
            "template_file": "格式模板",
            "output_file":   "輸出檔名",
            "output_type":   "輸出格式",
            "rows":          "資料筆數",
            "mapped":        "映射欄位",
            "total_cols":    "目標欄位數",
        })
        display_cols = [
            c for c in [
                "日期","時間","原始檔案","格式模板",
                "輸出檔名","輸出格式","資料筆數","映射欄位","目標欄位數",
            ] if c in hist_df.columns
        ]
        st.dataframe(hist_df[display_cols], use_container_width=True, hide_index=True)

        col_dl, col_clr, _ = st.columns([2, 2, 4])
        with col_dl:
            st.download_button(
                "⬇️ 匯出歷史 (JSON)",
                data=json.dumps(history, ensure_ascii=False, indent=2).encode("utf-8"),
                file_name=f"history_{datetime.now().strftime('%Y%m%d')}.json",
                mime="application/json",
            )
        with col_clr:
            if st.button("🗑️ 清除所有記錄", type="secondary"):
                save_history([])
                st.rerun()
